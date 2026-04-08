# -*- coding: utf-8 -*-
"""知识库加载服务"""

import os
import json
import re
from services.session_service import log_message
from config import UPLOAD_DIR, PLAYWRIGHT_DIR

# 全局知识库缓存
current_knowledge_content = ""
current_knowledge_file = ""


def load_knowledge_content(knowledge_file: str) -> str:
    """加载知识库文件"""
    global current_knowledge_content, current_knowledge_file
    
    if not knowledge_file:
        return ""
    
    # 如果传入的是列表，使用 load_multiple_knowledge 处理
    if isinstance(knowledge_file, list):
        return load_multiple_knowledge(knowledge_file)
    
    if knowledge_file == current_knowledge_file and current_knowledge_content:
        return current_knowledge_content
    
    filepath = os.path.join(UPLOAD_DIR, knowledge_file)
    if not os.path.exists(filepath):
        return ""
    
    try:
        ext = os.path.splitext(knowledge_file)[1].lower()
        
        if ext in (".txt", ".md"):
            with open(filepath, "r", encoding="utf-8") as f:
                content = f.read()
        elif ext == ".docx":
            from docx import Document
            doc = Document(filepath)
            content = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(filepath)
            content = "\n".join([
                " ".join([str(cell) if cell else "" for cell in row])
                for sheet in wb.worksheets for row in sheet.iter_rows(values_only=True)
            ])
        elif ext == ".pdf":
            import fitz
            doc = fitz.open(filepath)
            content = "\n".join([page.get_text() for page in doc])
            doc.close()
        else:
            with open(filepath, "r", encoding="utf-8") as f:
                content = f.read()
        
        if content:
            current_knowledge_content = content
            current_knowledge_file = knowledge_file
            log_message(f"已加载知识库: {knowledge_file} ({len(content)}字符)")
            return content
    except Exception as e:
        log_message(f"读取知识库失败: {e}", "ERROR")
    
    return ""


def load_multiple_knowledge(knowledge_files) -> str:
    """加载并合并多个知识库内容"""
    if not knowledge_files:
        return ""
    
    # 如果是单个字符串，转为列表处理
    if isinstance(knowledge_files, str):
        knowledge_files = [knowledge_files] if knowledge_files else []
    
    # 展平嵌套列表
    flat_files = []
    for f in knowledge_files:
        if isinstance(f, list):
            flat_files.extend(f)
        elif isinstance(f, str) and f:
            flat_files.append(f)
    
    if not flat_files:
        return ""
    
    contents = []
    for f in flat_files:
        if isinstance(f, str) and f:
            content = load_knowledge_content(f)
            if content:
                contents.append(f"=== {f} ===\n{content}")
    
    merged = "\n\n".join(contents)
    if merged:
        log_message(f"已合并 {len(contents)} 个知识库 ({len(merged)}字符)")
    return merged


def load_saved_questions(session_id: str = "", target: str = "") -> list:
    """从文件加载已保存的问题，支持多轮对话格式
    
    Args:
        session_id: 会话ID
        target: 目标类型 ("concurrent" 渠道Web测试, "solo" 单网站测试, "" 自动检测)
    
    Returns:
        问题列表
    """
    questions = []
    
    # 定义搜索路径列表（按优先级排序）
    search_paths = []
    
    if session_id:
        # 单网站测试的问题文件路径（优先）
        if target in ("solo", ""):
            search_paths.append({
                "path": os.path.join(PLAYWRIGHT_DIR, "solo_worker_PlayWright", "questions", session_id, "test_questions.txt"),
                "meta_path": os.path.join(PLAYWRIGHT_DIR, "solo_worker_PlayWright", "questions", session_id, "questions_meta.json"),
                "desc": "调试测试session问题文件"
            })
        # 渠道Web测试的问题文件路径
        if target in ("concurrent", ""):
            search_paths.append({
                "path": os.path.join(PLAYWRIGHT_DIR, "max_worker", "questions", session_id, "test_questions.txt"),
                "meta_path": os.path.join(PLAYWRIGHT_DIR, "max_worker", "questions", session_id, "questions_meta.json"),
                "desc": "渠道测试session问题文件"
            })
    
    # 默认问题文件路径
    if target in ("solo", ""):
        search_paths.append({
            "path": os.path.join(PLAYWRIGHT_DIR, "solo_worker_PlayWright", "test_questions.txt"),
            "meta_path": os.path.join(PLAYWRIGHT_DIR, "solo_worker_PlayWright", "questions_meta.json"),
            "desc": "调试测试默认问题文件"
        })
    if target in ("concurrent", ""):
        search_paths.append({
            "path": os.path.join(PLAYWRIGHT_DIR, "max_worker", "test_questions.txt"),
            "meta_path": os.path.join(PLAYWRIGHT_DIR, "max_worker", "questions_meta.json"),
            "desc": "渠道测试默认问题文件"
        })
    
    # 按优先级搜索
    for sp in search_paths:
        path = sp["path"]
        meta_path = sp["meta_path"]
        desc = sp["desc"]
        
        if os.path.exists(path):
            try:
                with open(path, "r", encoding="utf-8") as f:
                    content = f.read()
                
                # 检查是否有元数据文件
                meta_data = None
                is_multi_turn = False
                if os.path.exists(meta_path):
                    try:
                        with open(meta_path, "r", encoding="utf-8") as mf:
                            meta_data = json.load(mf)
                            # 检查是否是列表格式（每个问题的元数据）
                            if isinstance(meta_data, list) and len(meta_data) > 0:
                                # 新格式：列表中每个元素是一个问题的元数据
                                # 只有多个组才是多轮对话
                                group_indices = set(item.get('group_index', 0) for item in meta_data if isinstance(item, dict))
                                is_multi_turn = len(group_indices) > 1
                            else:
                                # 旧格式
                                is_multi_turn = meta_data.get("is_multi_turn", False)
                    except:
                        pass
                
                # 如果有元数据，使用元数据构建问题列表
                if meta_data and isinstance(meta_data, list) and len(meta_data) > 0:
                    # 按组重新组织问题
                    groups = {}
                    for item in meta_data:
                        if isinstance(item, dict) and 'question' in item:
                            idx = item.get('group_index', 0)
                            if idx not in groups:
                                groups[idx] = []
                            groups[idx].append({
                                'question': item.get('question', ''),
                                'question_type': item.get('question_type', 'normal'),
                                'group_index': idx
                            })
                    
                    if len(groups) > 1 or (len(groups) == 1 and len(groups.get(0, [])) > 1):
                        # 多轮对话格式
                        questions = [groups[k] for k in sorted(groups.keys())]
                    else:
                        # 单轮对话格式
                        questions = groups.get(0, [])
                    
                    log_message(f"从{desc}加载了{len(questions)}个问题（含类型信息）", session_id=session_id)
                    return questions
                
                # 没有元数据，使用原有逻辑
                if is_multi_turn:
                    # 多轮对话格式：用空行分隔组
                    groups = re.split(r'\n\s*\n', content.strip())
                    questions = []
                    for group in groups:
                        group_questions = [l.strip() for l in group.split("\n") if l.strip() and not l.startswith("#")]
                        if group_questions:
                            questions.append(group_questions)
                else:
                    # 单轮对话格式
                    questions = [l.strip() for l in content.split("\n") if l.strip() and not l.startswith("#")]
                
                log_message(f"从{desc}加载了{len(questions)}个问题", session_id=session_id)
                return questions
            except Exception as e:
                log_message(f"读取{desc}失败: {e}", session_id=session_id)
    
    return questions
